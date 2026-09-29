#!/usr/bin/env python3
"""调用大模型接口，为任务生成多条不重复的每日工作描述。

接口地址和密钥不写死在代码里，通过环境变量自行配置：
  OLLAMA_URL       生成接口地址，支持两种风格（按 URL 自动判断）：
                   - Ollama 原生接口，例如 https://your-ollama-host:port/api/generate
                   - OpenAI 兼容接口（DeepSeek、vLLM 等），
                     例如 https://api.deepseek.com/chat/completions
  OLLAMA_API_KEY   接口的 Bearer token
  OLLAMA_MODEL     模型名（可选。不设时按接口风格自动选：Ollama 风格默认
                   qwen3.6:35b-a3b，OpenAI 风格默认 deepseek-chat）

  WORK_LOG         真实工作日志文件路径（可选）。默认自动探测：优先
                   work_log.xlsx（Excel 两列表格：第一列日期段如 9.1-9.12，
                   第二列这段时间做了什么，行数随意增删），其次 work_log.md。
                   有实质内容时注入提示词当背景素材，大模型结合"哪段时间干
                   了什么"为每天生成贴合实际工作的描述。跑完 fill_progress.py
                   且所有任务的描述都生成完毕后，日志自动归档到 *_archive
                   （xlsx 追加带时间戳的 sheet / md 追加到文末），原文件重置
                   回空模板，下月直接填新行。

不配置也能正常跑：会跳过大模型调用，直接用模板生成的兜底描述，不会报错
卡住整个流程。
"""

import datetime
import os
import re

import openpyxl
import requests

API_KEY = os.getenv("OLLAMA_API_KEY", "")
URL = os.getenv("OLLAMA_URL", "")
MODEL = os.getenv("OLLAMA_MODEL", "")  # 留空时按接口风格自动选默认模型

# 工作日志表格的表头词（判断"这一行是不是表头"用），xlsx 和 md 共用
_HEADER_WORDS = {"日期", "日期段", "时间", "工作内容", "内容"}

# 空模板：归档重置时写回 work_log.md，用户每次在表格里填本月内容即可。
# 说明全放在 HTML 注释里，_extract_substance 会把注释整体剔除，模板说明
# 和示例行永远不会混进发给大模型的素材。
WORK_LOG_TEMPLATE = """<!-- 工作日志模板：每次填报前，把本月做了什么填进下面表格，行数随意增删。
  第一列写日期段，如 9.1-9.12 或 9.22-（表示 9.22 到月底）；
  第二列写这段时间做了什么，一两句话即可。
  示例行：| 9.1-9.12 | 日志回访功能开发：回访记录列表、批量导入、失败重试 |
  跑 fill_progress.py 并生成完全部任务的描述后，本文件会自动归档到
  work_log_archive.md 并重置回本模板，下个月直接填新行即可。
  注意：内容会原文发送给你配置的大模型接口，不要写敏感信息。 -->

| 日期段 | 工作内容 |
|---|---|
|  |  |
|  |  |
|  |  |
"""

HEADERS = {
    "Authorization": f"Bearer {API_KEY}",
    "Content-Type": "application/json",
}


def configure_llm(
    url: str | None = None,
    api_key: str | None = None,
    model: str | None = None,
) -> None:
    """运行时更新接口配置（Web 控制台等外部入口用）。

    None 表示该项保持不变；空串表示恢复默认——模型名回到"按接口风格自动选"，
    URL/Key 回到未配置状态（大模型调用被跳过，走模板兜底）。"""
    global URL, API_KEY, MODEL, HEADERS
    if url is not None:
        URL = url
    if api_key is not None:
        API_KEY = api_key
        HEADERS = {"Authorization": f"Bearer {API_KEY}", "Content-Type": "application/json"}
    if model is not None:
        MODEL = model


def _llm_configured() -> bool:
    return bool(URL and API_KEY)


def _is_ollama_style() -> bool:
    """按 URL 区分接口风格：路径含 /api/generate 的按 Ollama 原生格式请求
    （prompt 字段进、response 字段出）；其余（DeepSeek、vLLM、Ollama 的
    OpenAI 兼容端点等）按 OpenAI chat/completions 格式请求（messages 进、
    choices[0].message.content 出）。"""
    return "/api/generate" in URL


def _default_model() -> str:
    return "qwen3.6:35b-a3b" if _is_ollama_style() else "deepseek-chat"


def _call_llm(prompt: str) -> str:
    model = MODEL or _default_model()

    if _is_ollama_style():
        resp = requests.post(
            URL,
            headers=HEADERS,
            json={"model": model, "prompt": prompt, "stream": False, "think": False},
            timeout=120,
        )
        resp.raise_for_status()
        return resp.json().get("response", "").strip()

    resp = requests.post(
        URL,
        headers=HEADERS,
        json={
            "model": model,
            "messages": [{"role": "user", "content": prompt}],
            "stream": False,
        },
        timeout=120,
    )
    resp.raise_for_status()
    return resp.json()["choices"][0]["message"]["content"].strip()


def _work_log_candidates() -> list[str]:
    """日志文件候选路径，按优先级排列：环境变量 WORK_LOG 指定则只用它；
    否则 Excel（work_log.xlsx）优先、Markdown（work_log.md）兜底。"""
    env = os.getenv("WORK_LOG", "")
    if env:
        return [env]
    return ["work_log.xlsx", "work_log.md"]


def _read_text(path: str) -> str:
    """读文本文件原文（不做任何过滤），不存在、不可读返回空串。
    Windows 记事本旧版本可能把文件存成 GBK，UTF-8 解不开时再退一步用 GBK 读。"""
    try:
        with open(path, "r", encoding="utf-8") as f:
            return f.read()
    except UnicodeDecodeError:
        try:
            with open(path, "r", encoding="gbk") as f:
                return f.read()
        except (OSError, UnicodeDecodeError):
            return ""
    except OSError:
        return ""


def _cell_text(cell) -> str:
    """把 Excel 单元格统一转成干净的文本。日期类型的单元格（Excel 会把
    2026/9/1 这类输入自动识别成日期）格式化成 YYYY-MM-DD，别让 str() 吐出
    "2026-09-01 00:00:00" 这种带零点的长串。"""
    if cell is None:
        return ""
    if isinstance(cell, datetime.datetime):
        if (cell.hour, cell.minute, cell.second) == (0, 0, 0):
            return cell.strftime("%Y-%m-%d")
        return cell.strftime("%Y-%m-%d %H:%M")
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return str(cell).strip()


def _read_xlsx_rows(path: str) -> list[list[str]]:
    """读 Excel 第一个工作表，返回填了内容的数据行（剔除表头行/空行），
    每行是干净的单元格文本列表。文件不存在或读不了返回空列表。"""
    if not os.path.exists(path):
        return []
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return []
    try:
        rows: list[list[str]] = []
        for row in wb.worksheets[0].iter_rows(values_only=True):
            cells = [_cell_text(c) for c in row]
            while cells and not cells[-1]:
                cells.pop()  # 去掉行尾空单元格
            if not cells:
                continue
            nonempty = [c for c in cells if c]
            if nonempty and all(c in _HEADER_WORDS for c in nonempty):
                continue  # 表头行
            rows.append(cells)
        return rows
    finally:
        wb.close()


def _extract_substance(text: str) -> list[str]:
    """提取日志里的实质内容行：剔除 HTML 注释块（模板说明）、Markdown 表格的
    表头/分隔线/空数据行。填了内容的表格行、自由文本行（如月份标题）都保留。

    用途有二：判断"模板还没填过"（返回空列表），以及只把干净的内容行发给
    大模型（模板注释里的示例行不会混进素材）。"""
    header_words = _HEADER_WORDS
    rows: list[str] = []
    in_comment = False
    for line in text.splitlines():
        s = line.strip()
        if not s:
            continue
        if in_comment:
            if "-->" in s:
                in_comment = False
            continue
        if "<!--" in s:
            # 单行注释 <!-- ... --> 直接跳过；注释没闭合则进入跨行注释状态
            if "-->" not in s.split("<!--", 1)[1]:
                in_comment = True
            continue
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            meaningful = [
                c
                for c in cells
                if c and c not in header_words and not set(c) <= set("-: ")
            ]
            if not meaningful:
                continue  # 表头 / 分隔线 / 空数据行
            rows.append(s)
        else:
            rows.append(s)
    return rows


def _work_log_substance(path: str) -> list[str]:
    """单个日志文件的实质内容：xlsx 返回"日期段 | 内容"拼接行，文本格式返回
    剔除注释/表头/空行后的行。文件不存在或只是空模板返回空列表。"""
    if path.lower().endswith(".xlsx"):
        return [" | ".join(cells) for cells in _read_xlsx_rows(path)]
    raw = _read_text(path)
    return _extract_substance(raw) if raw else []


def _load_work_log() -> str:
    """按优先级探测日志文件，返回第一个有实质内容文件的干净内容。

    所有候选都不存在、不可读、或只是还没填过的空模板，都返回空串，行为退回
    "仅按任务名生成"，不影响主流程。"""
    for path in _work_log_candidates():
        rows = _work_log_substance(path)
        if rows:
            return "\n".join(rows)
    return ""


def _write_text_template(path: str) -> None:
    with open(path, "w", encoding="utf-8") as f:
        f.write(WORK_LOG_TEMPLATE)


def _write_xlsx_template(path: str) -> None:
    """写一个空的 Excel 日志模板：两列表头 + 3 个空行，列宽调好方便填。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "工作日志"
    ws.append(["日期段", "工作内容"])
    for _ in range(3):
        ws.append(["", ""])
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 60
    wb.save(path)


def _archive_xlsx(path: str) -> str:
    """把 Excel 日志的数据行归档：追加一个以归档时间命名的 sheet 到
    *_archive.xlsx（没有则创建）。返回归档文件路径。"""
    archive_path = os.path.splitext(path)[0] + "_archive.xlsx"
    stamp = datetime.datetime.now().strftime("%Y-%m-%d %H-%M")
    if os.path.exists(archive_path):
        wb = openpyxl.load_workbook(archive_path)
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # 删掉新建工作簿自带的空 sheet
    title = stamp[:31]  # Excel sheet 名最长 31 字符
    while title in wb.sheetnames:
        title += "_"
    ws = wb.create_sheet(title=title)
    ws.append(["日期段", "工作内容"])
    for cells in _read_xlsx_rows(path):
        ws.append(cells)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 60
    wb.save(archive_path)
    return archive_path


def archive_and_reset_work_log() -> str | None:
    """把当前有实质内容的工作日志归档，再把该文件重置回空模板。

    按优先级找到第一个"填了内容"的日志文件才动作；所有候选都为空（还没填）
    时什么都不动、返回 None。归档是追加写（xlsx 追加新 sheet / md 追加到
    文末），每次运行的历史逐次累积、不会丢。返回归档文件路径。

    归档/重置遇文件被占用（Excel 开着、同步盘锁定）不打断主流程——填报
    都已提交成功，收尾这步失败不该把整个运行标成"出错"。打印警告、保留
    日志原文件，下次运行自动重试。"""
    for path in _work_log_candidates():
        if not _work_log_substance(path):
            continue

        try:
            if path.lower().endswith(".xlsx"):
                archive_path = _archive_xlsx(path)
                reset = _write_xlsx_template
            else:
                archive_path = os.path.splitext(path)[0] + "_archive.md"
                stamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
                with open(archive_path, "a", encoding="utf-8") as f:
                    f.write(f"\n<!-- 归档于 {stamp}（所有任务描述生成完毕后自动归档） -->\n")
                    f.write(_read_text(path).rstrip() + "\n")
                reset = _write_text_template
        except OSError as e:
            print(
                f"    [工作日志] 归档失败：{e}。多半是 {path} 或归档文件正被 "
                f"Excel 占用——关掉它们即可，日志内容原样保留，下次运行自动重试归档",
                flush=True,
            )
            return None

        try:
            reset(path)
        except OSError as e:
            print(
                f"    [工作日志] 归档成功但重置模板失败：{e}。关闭 {path} 后重新"
                f"运行会自动重试（归档文件里会多一份重复记录，手动删掉多余 sheet 即可）",
                flush=True,
            )
        return archive_path
    return None


# 行首日期前缀（"2026-09-01"、"09-01"、"9.1"、"9月1日"、"9.1-9.12"、"第1天"等）。
# 日期只是提示词里给大模型对齐用的素材，RDM 的工作描述里不能出现日期字样，
# 所以除了 prompt 里要求，生成结果再统一剥一遍行首日期兜底。
_DATE_PREFIX_RE = re.compile(
    r"""^(?:
        (?:\d{4}\s*[-/.年]\s*)?                    # 可选年份
        \d{1,2}\s*[-/.月]\s*\d{1,2}                # 月-日：9-1 / 09.01 / 9月1
        (?:\s*[日号])?                             # 可选"日/号"
        (?:\s*[-–—~至]\s*                          # 可选日期段终点：9.1-9.12
            (?:\d{4}\s*[-/.年]\s*)?\d{1,2}(?:\s*[-/.月]\s*\d{1,2})?(?:\s*[日号])?)?
        (?:\s*[：:、|，,]\s*|\s+)                   # 日期后必须跟分隔符或空格
        |
        第\s*\d+\s*天\s*[：:、|，,]?\s*             # "第1天：..."
    )""",
    re.VERBOSE,
)


def _clean_lines(text: str) -> list[str]:
    lines = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        # 去掉行首日期前缀（"09-01 梳理..." -> "梳理..."）。必须放在编号前缀
        # 之前处理，否则 "09-01" 会先被编号规则截成 "-01"
        line = _DATE_PREFIX_RE.sub("", line, count=1).strip()
        # 去掉常见的编号前缀，如 "1. " "1、" "- " "* "；行首形如 "9.1"（数字+
        # 分隔符+数字，是日期或版本号）的不当作编号，避免截成 "1版本..."
        if not re.match(r"^\d{1,4}\s*[-/.]\s*\d", line):
            line = re.sub(r"^[\d一二三四五六七八九十]+[.、\)\s]*", "", line)
        line = line.lstrip("-*• ").strip()
        if line:
            lines.append(line)
    return lines


def generate_daily_descriptions(
    task_name: str, n: int, dates: list[str] | None = None
) -> list[str]:
    """为任务生成 n 条互不相同的一句话工作描述，按开发流程不同阶段递进。

    dates 是这 n 天对应的执行日期（如 ["2026-09-01", ...]）。配合
    work_log.xlsx / work_log.md 里"哪段时间干了什么"的记录，大模型能把每天
    的描述对齐到当天实际在做的事，而不是纯靠任务名脑补通用剧情。
    没配置 OLLAMA_URL/OLLAMA_API_KEY，或者请求失败，都会直接用模板兜底，
    不会抛异常打断整个填报流程。"""
    if n <= 0:
        return []

    results: list[str] = []

    if not _llm_configured():
        print("    [大模型] 未配置 OLLAMA_URL / OLLAMA_API_KEY，跳过大模型调用，使用模板描述", flush=True)
    else:
        work_log = _load_work_log()
        if work_log:
            print("    [大模型] 已读取工作日志（work_log.xlsx / work_log.md，可用 WORK_LOG 指定路径），结合真实工作日志生成描述", flush=True)
        else:
            print("    [大模型] 工作日志还没填（work_log.xlsx / work_log.md 都没有实质内容），仅按任务名生成通用描述", flush=True)

        attempts = 0
        while len(results) < n and attempts < 4:
            attempts += 1
            remaining = n - len(results)
            avoid = ""
            if results:
                avoid_list = "\n".join(f"- {d}" for d in results)
                avoid = f"\n以下描述已经用过，不要重复或过于相似：\n{avoid_list}\n"

            if work_log:
                log_hint = (
                    f"\n以下是我的真实工作日志（自由格式，哪段时间在做什么以它为准）：\n"
                    f"{work_log}\n"
                )
                stage_hint = (
                    "内容必须贴合日志里该日期附近实际在做的事，不要照抄原文、"
                    "改写成自然的一句话日报；日志没直接覆盖的日期写推进该任务相关工作的表述，"
                    "但不要与日志矛盾"
                )
            else:
                log_hint = ""
                stage_hint = (
                    "内容围绕需求分析、方案设计、编码实现、联调、测试、修复问题、"
                    "文档整理等不同阶段展开"
                )

            date_hint = ""
            if dates:
                rest_dates = dates[len(results):]
                if rest_dates:
                    date_hint = (
                        "\n下面是需要生成描述的日期，按顺序每条输出对应一天的描述：\n"
                        + "\n".join(rest_dates) + "\n"
                    )

            prompt = (
                f"你是一名研发工程师，正在为任务《{task_name}》填写连续多天的每日工作进展描述。\n"
                f"{log_hint}"
                f"请生成 {remaining} 条不同的中文工作描述，每条一句话（20~40字），语气自然、具体，"
                f"像真实的研发日报，{stage_hint}，条目之间不要重复或雷同。{date_hint}{avoid}\n"
                f"严格按下面格式输出，每条一行，只写当天的工作内容本身，行首不要带任何日期"
                f"（如 2026-09-01、09-01、9月1日、第1天，系统会单独记录日期），"
                f"不要照抄日志里的日期，不要加编号、不要加多余说明：\n"
            )

            print(f"    [大模型] 第 {attempts} 次请求，还差 {remaining} 条...", flush=True)
            try:
                text = _call_llm(prompt)
            except Exception as e:
                print(f"    [大模型] 请求失败（{e}），改用模板兜底", flush=True)
                break
            for line in _clean_lines(text):
                if line not in results:
                    results.append(line)
                if len(results) >= n:
                    break
            print(f"    [大模型] 已生成 {len(results)}/{n} 条", flush=True)

    if len(results) < n:
        # 兜底：不足的部分用简单模板补齐，保证流程不中断
        for i in range(len(results), n):
            results.append(f"继续推进《{task_name}》相关工作，处理第 {i + 1} 阶段的开发与验证。")

    return results[:n]


if __name__ == "__main__":
    for d in generate_daily_descriptions(
        "统一检索功能优化", 3, ["2026-09-01", "2026-09-02", "2026-09-03"]
    ):
        print("-", d)
